"""Highlight the query optimizations that are in the current Analytics Advisor code.

Writes docs/Query_Optimization_Applied.docx.
Does not replace the architecture deck, the overview, or the older preview-path note.
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "Query_Optimization_Applied.docx"
NAVY = RGBColor(0x1E, 0x27, 0x61)


def heading(doc: Document, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = NAVY


def para(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)


def bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text, style="List Bullet")
    for run in p.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(11)


def mono(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)


def table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    grid = doc.add_table(rows=1 + len(rows), cols=len(headers))
    grid.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = grid.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(10)
        run.font.color.rgb = NAVY
    for ri, row in enumerate(rows):
        for ci, value in enumerate(row):
            cell = grid.rows[ri + 1].cells[ci]
            cell.text = ""
            run = cell.paragraphs[0].add_run(value)
            run.font.name = "Calibri"
            run.font.size = Pt(10)
    doc.add_paragraph()


def build() -> None:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    title = doc.add_paragraph()
    run = title.add_run("Query optimization applied in Analytics Advisor")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(22)
    run.font.color.rgb = NAVY

    sub = doc.add_paragraph()
    srun = sub.add_run(
        "What the current code does, and why Save, Open, and a wide result no longer fail the same way.  "
        f"{date.today():%d %B %Y}."
    )
    srun.font.name = "Calibri"
    srun.font.size = Pt(11)
    srun.italic = True
    srun.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    para(
        doc,
        "This note describes the optimizations that are in the application code on master "
        "(commit a34b7e7, https://github.com/kp-jpr/analytics-advisor). "
        "A further speed-up for the Run estimate is written as a plan and is not in this build.",
    )

    heading(doc, "1. The problem the code was solving", 1)
    para(
        doc,
        "An officer builds a match in Query Builder: a source table, a destination table, "
        "optional extra tables, filters, and the columns to show. Three actions used to share one expensive step.",
    )
    bullet(doc, "Save stored the query definition, but first counted rows in the lakehouse, then counted them again before the response returned.")
    bullet(doc, "Open loaded that definition only after the same live count. A count that failed, or a count of a numeric column passed to lower(), meant the query never opened.")
    bullet(doc, "Run still counts before the join. That count is the guard that refuses a match estimated above 50 million rows. That guard stays on Run.")
    para(doc, "A filtered count was the statement officers already saw take a long time:")
    mono(doc, "SELECT count(*) FROM (SELECT * FROM silver_data.shaladarpan_txn.tbl_txn_studentdata t WHERE t.dob = …)")
    para(doc, "A fuzzy estimate on a numeric identifier failed before the match started:")
    mono(
        doc,
        "SELECT CAST(approx_distinct(substr(lower(golden_data.smart_services."
        "tbl_smart_citizen_address.jan_member_id), 1, 3)) AS BIGINT) FROM …",
    )
    para(
        doc,
        "Presto accepts lower() only on text. jan_member_id is a bigint, so the estimate stopped with "
        "Unexpected parameters (bigint) for function lower. Selecting more than eight output columns "
        "stopped earlier still, with sourceDisplayColumns must have 0 to 8 entries, shown in red as "
        "Analysis service error 400.",
    )

    heading(doc, "2. What the current code optimizes", 1)
    table(
        doc,
        ["Action", "What the code does now", "Lakehouse row count"],
        [
            ["Save", "Checks tables, columns, join shape, and filters, then stores the row once.", "Not run"],
            ["Open", "Same shape check, then returns the stored definition so Extract can fill in.", "Not run"],
            [
                "Run, preview, export",
                "Full plan, including the 50 million row ceiling.",
                "Still run when statistics are missing, a filter applies, or the join is fuzzy",
            ],
            [
                "Output columns",
                "Every selected attribute is accepted. Columns must still belong to that side's table.",
                "Not a reason to reject the query",
            ],
        ],
    )

    heading(doc, "3. Save and Open skip the live count", 1)
    para(
        doc,
        "RecordMatchService.planMatch is the full plan. For a join it calls enforceEstimatedRowCeiling, "
        "which issues count(*) and, when needed, approx_distinct against the lakehouse. "
        "planMatchForSave runs the same validation and builds the same SQL shape, and it does not call that ceiling.",
    )
    para(doc, "SavedQueryService uses planMatchForSave in three places, once each:")
    bullet(doc, "create validates, inserts the saved_query row, and returns that row. It does not call getForCurrentUser, which used to plan the query a second time.")
    bullet(doc, "update does the same single validation, then returns the row just written.")
    bullet(doc, "getForCurrentUser is Open. It validates under the current officer and returns the stored request. A shared query is still checked for tables and columns the officer can use. It is not counted before it is shown.")
    para(
        doc,
        "The 50 million row ceiling remains a guard for running a match. Storing or reopening a definition does not scan the lakehouse to compute it. "
        "When the officer presses Run, planMatch still refuses an estimate above 50 million rows.",
    )

    heading(doc, "4. Check the column type, then cast", 1)
    para(
        doc,
        "The compiler reads the live column type before it writes a filter or a join. "
        "SqlTypeFamily groups those types as text, number, boolean, temporal, or unknown. "
        "Presto already coerces values inside one family (integer with bigint, varchar lengths, date with timestamp). "
        "It does not coerce across families, which is why date = varchar and lower(bigint) failed.",
    )
    para(doc, "The default comparison follows the column:")
    bullet(doc, "A date or timestamp column stays a date. The officer's text is bound as CAST(? AS DATE), CAST(? AS TIMESTAMP), or the matching time type. A value that is not a date still fails the query, so a bad date is not silently dropped.")
    bullet(doc, "A number column compared with text uses TRY_CAST of the text side to DOUBLE. A Java number compared with a text column still casts the column to DOUBLE, so 123 and 0123 can meet.")
    bullet(doc, "A date joined to text is compared as a date: TRY_CAST of the text side to DATE, and the date column left as a date.")
    bullet(doc, "An explicit admin setting of text, number, or date still wins over this default.")
    para(
        doc,
        "Fuzzy matching is a text measure. Both sides are cast with TypeCoercion.asText before lower(), substr(), and levenshtein_distance. "
        "A bigint such as jan_member_id becomes CAST(column AS VARCHAR) and only then lower(). "
        "The row-count estimate uses that same cast, so the approx_distinct blocking key does not call lower() on a bigint.",
    )
    para(
        doc,
        "Direct databases do not speak every Presto cast. AnalysisExecutionRouter rewrites the temporal casts after the query is routed. "
        "SQL Server uses DATETIME2 and DATETIMEOFFSET. Oracle uses TO_DATE and TO_TIMESTAMP. MySQL and MariaDB use DATETIME. "
        "DB2 maps timestamp with time zone to TIMESTAMP. PostgreSQL and the lakehouse keep the Presto form.",
    )

    heading(doc, "5. A handled failure is a notice, not a red service error", 1)
    para(
        doc,
        "A comparison the database cannot apply is returned as HTTP 400 with code TYPE_MISMATCH. "
        "The officer sees one sentence: enter a date as YYYY-MM-DD, and each engine compares a date only with a date. "
        "The failing SQL is not appended to that sentence. Other rejected queries stay QUERY_FAILED and still include the failing SQL. "
        "A down analytical engine stays a connectivity message: The analytical engine is unavailable. Check the Presto connection, then try again.",
    )
    para(
        doc,
        "The page no longer prefixes a handled message with Analysis service error 400. "
        "presentQueryError strips that prefix. Save, Open, preview, and the match result show the sentence in a notice "
        "(a left border in the warning colour), not in red error text.",
    )

    heading(doc, "6. Output columns are not capped at eight", 1)
    para(
        doc,
        "validateDisplayColumns still requires each output column to belong to that side's table. "
        "It no longer rejects a list longer than eight. The same is true for display columns on an additional table. "
        "Select all on a wide table is therefore a valid query.",
    )
    para(
        doc,
        "Join groups and comparison groups are still limited to eight. That limit is about the join shape, not about how many attributes are shown.",
    )

    heading(doc, "7. What Run still does, and the next plan", 1)
    para(
        doc,
        "Run, SQL preview, and export still call planMatch. When catalog statistics are missing, a filter is present, "
        "or the join is fuzzy, planning still asks the lakehouse for counts before the match SQL starts. "
        "A three-table filtered match has already spent about 24 seconds in that estimate and about 10 seconds in the join. "
        "The estimate queries run one after another, and a filtered count is still wrapped as SELECT count(*) FROM (SELECT * FROM table …).",
    )
    para(doc, "The written next step, not yet in this build, keeps the 50 million ceiling and changes only how the estimate is collected:")
    bullet(doc, "Count a filtered table directly: SELECT count(*) FROM table t WHERE …, and the same direct FROM for approx_distinct. The SELECT * wrapper goes away. The filter text and the bound values stay the same.")
    bullet(doc, "Run the independent counts together on a pool of four threads, then apply the same fan-out formula.")
    bullet(doc, "Remember SHOW STATS for ten minutes per table. A filtered count is not cached, because the filter changes the answer.")
    para(
        doc,
        "Single-source extract has no join fan-out, so it already skips this ceiling. Save and Open stay on planMatchForSave.",
    )

    heading(doc, "8. Where to look in the code", 1)
    table(
        doc,
        ["Behaviour", "Code"],
        [
            ["Full plan and the 50 million ceiling", "RecordMatchService.planMatch, enforceEstimatedRowCeiling"],
            ["Save and Open plan, no live count", "RecordMatchService.planMatchForSave, SavedQueryService.validateForSave"],
            ["Casts for join, filter, and fuzzy text", "TypeCoercion, RuleCompiler, sideDistinctEstimate"],
            ["Casts rewritten per direct database", "AnalysisExecutionRouter.adaptTemporalCasts"],
            ["Type-mismatch sentence, HTTP 400", "QueryFailureMessages, DataAccessExceptionHandler"],
            ["Notice instead of a red service error", "queryErrorText.presentQueryError, srse-query-notice"],
            ["No eight-column display cap", "RecordMatchService.validateDisplayColumns"],
        ],
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
