"""Build docs/Query_Optimization_Exception_Handling_and_Standards.docx."""
from __future__ import annotations

from datetime import date
from pathlib import Path

import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch

ROOT = Path(__file__).resolve().parent.parent
ASSETS = ROOT / "tech_doc_assets"
OUT = ROOT / "Query_Optimization_Exception_Handling_and_Standards.docx"
NAVY = "#1E2761"


def _box(ax, x, y, w, h, text, fc="#FFFFFF", fs=8):
    ax.add_patch(
        FancyBboxPatch(
            (x, y),
            w,
            h,
            boxstyle="round,pad=0.02,rounding_size=0.08",
            linewidth=1.15,
            edgecolor=NAVY,
            facecolor=fc,
        )
    )
    ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=fs, color="#1A1A1A")


def _arrow(ax, x1, y1, x2, y2):
    ax.add_patch(
        FancyArrowPatch(
            (x1, y1),
            (x2, y2),
            arrowstyle="-|>",
            mutation_scale=12,
            linewidth=1.15,
            color="#444444",
        )
    )


def draw_diagram() -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    fig, ax = plt.subplots(figsize=(11.2, 7.2))
    ax.set_xlim(0, 11.2)
    ax.set_ylim(0, 7.2)
    ax.axis("off")
    ax.set_title(
        "Current SQL preview path and the proposed path",
        fontsize=13,
        fontweight="bold",
        color=NAVY,
        pad=8,
    )
    ax.text(2.7, 6.55, "Current", ha="center", fontsize=11, fontweight="bold", color=NAVY)
    ax.text(8.4, 6.55, "Proposed", ha="center", fontsize=11, fontweight="bold", color=NAVY)

    current = [
        (5.55, "Query Builder edit", "#CADCFC"),
        (4.35, "POST /api/analysis/match.sql\nevery 400 ms", "#FFFFFF"),
        (3.05, "validateColumn once\nper display column", "#FFF4E5"),
        (1.55, "information_schema walk\ncatalogs, schemas, tables, columns", "#FDECEC"),
        (0.25, "Presto connection to DB2\nheld for minutes", "#FDECEC"),
    ]
    proposed = [
        (5.55, "Stable cascade\ncached metadata", "#CADCFC"),
        (4.35, "Preview only after the\nquestion stops changing", "#FFFFFF"),
        (3.05, "validateColumns once\nper table", "#EAF6EA"),
        (1.55, "Registered columns from cache\nno live catalog walk", "#EAF6EA"),
        (0.25, "Match runs on direct JDBC\nor Presto only when federated", "#EAF6EA"),
    ]
    for y, text, fc in current:
        _box(ax, 0.35, y, 4.7, 0.95, text, fc=fc)
    for y, text, fc in proposed:
        _box(ax, 6.05, y, 4.7, 0.95, text, fc=fc)
    for y, _text, _fc in current[:-1]:
        _arrow(ax, 2.7, y, 2.7, y - 0.22)
    for y, _text, _fc in proposed[:-1]:
        _arrow(ax, 8.4, y, 8.4, y - 0.22)
    fig.savefig(ASSETS / "query_opt_paths.png", dpi=160, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def para(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)


def bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text, style="List Bullet")
    for run in p.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(11)


def table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    grid = doc.add_table(rows=1 + len(rows), cols=len(headers))
    grid.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = grid.rows[0].cells[i]
        cell.text = header
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.name = "Calibri"
                run.font.size = Pt(9)
    for ri, row in enumerate(rows):
        for ci, value in enumerate(row):
            cell = grid.rows[ri + 1].cells[ci]
            cell.text = value
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(9)
    doc.add_paragraph()


def build() -> None:
    draw_diagram()
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Analytics Advisor")
    run.bold = True
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor(0x1E, 0x27, 0x61)
    run.font.name = "Calibri"

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("Query Optimization, Exception Handling,\nand Code Standardization")
    run.font.size = Pt(18)
    run.font.name = "Calibri"

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(
        "Plan based on the current source and the 14 October 2026 deployment logs\n"
        f"Document date: {date.today().strftime('%d %B %Y')}\n"
        "This document records the plan. It does not change application code."
    )
    run.font.size = Pt(11)
    run.font.name = "Calibri"
    doc.add_page_break()

    heading(doc, "1. Purpose")
    para(
        doc,
        "This plan sets the order of work to make Query Builder faster, to make failures "
        "readable and consistent, and to keep later changes on the patterns the code already "
        "uses. The evidence is the deployment log from 14 October 2026, 12:36 to 12:44 UTC, "
        "and the current backend and frontend.",
    )
    para(
        doc,
        "In that window an officer used Query Builder against the Presto catalog db2_catalog, "
        "schema sws. Nginx recorded 42 catalog lists, 21 schema lists, and 21 table lists in "
        "about 30 seconds. Every one of those calls returned HTTP 200. The backend then held "
        "twelve Presto connections for about five to six minutes while it validated columns "
        "for SQL preview and one streamed match. Hikari reported each hold as a connection "
        "leak. All twelve connections were returned to the pool. PostgreSQL only wrote a "
        "normal checkpoint. The slowness is metadata traffic through Presto to DB2, not a "
        "failed login or a down operational database.",
    )

    heading(doc, "2. What is slow today")
    para(
        doc,
        "POST /api/analysis/match.sql is documented as a preview that returns SQL text and "
        "does not run the match. Before it returns that text, RecordMatchService validates "
        "every display column. LakehouseRegistryService.validateColumns is written to walk "
        "the catalog once per table. RecordMatchService.validateDisplayColumns does not use "
        "that batch method. It calls registry.validateColumn once per display column, around "
        "line 1770. The same single-column call remains for criteria and dedup columns.",
    )
    para(
        doc,
        "Each validateColumn call reaches LakehouseBrowseService.listColumns. Listing columns "
        "first calls validateTable, which lists every table in the schema from "
        "information_schema.tables. Validating the schema lists schemata, and validating the "
        "catalog lists catalogs. One column check therefore repeats the whole ladder. Against "
        "db2_catalog, Presto forwards that ladder to DB2 and keeps the JDBC connection until "
        "it finishes. Eleven preview threads and one match stream did this together, which is "
        "the twelve leak warnings in the log.",
    )
    para(
        doc,
        "The browser adds load of its own. LakehouseCascade refetches catalogs whenever "
        "browseFilter, value.catalog, or onChange changes, and refetches schemas and tables "
        "on the same pattern. A parent render that supplies a new onChange function, or an "
        "onChange that clears a selection the server still returns, starts the list calls "
        "again. ExtractRecordsTab schedules POST /api/analysis/match.sql 400 milliseconds "
        "after any change to the whole extract object. The previous request is marked "
        "cancelled in React state, but the HTTP call is not aborted, so the backend keeps "
        "the Presto connection until the metadata query ends.",
    )
    doc.add_picture(str(ASSETS / "query_opt_paths.png"), width=Inches(6.5))
    cap = doc.add_paragraph(
        "Figure 1. Left: each edit walks live DB2 metadata through Presto. "
        "Right: one validation per table, then cached columns and a short preview."
    )
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap.runs:
        run.italic = True
        run.font.name = "Calibri"
        run.font.size = Pt(10)
    doc.add_paragraph()

    heading(doc, "3. Query optimization plan")
    para(doc, "Do the items in this order. Each one removes work the log already showed.")

    heading(doc, "3.1 Validate a table's columns in one walk", 2)
    para(
        doc,
        "Inside RecordMatchService.validateDisplayColumns, check that every display column "
        "belongs to the side table, then call registry.validateColumns once with the full "
        "column list. Do the same for the criteria loop and the dedup column, grouping columns "
        "that share a table. Leave validateColumn for a caller that truly has one column. "
        "LakehouseRegistryService already says this batch exists for latency: resolving a "
        "column list walks SHOW CATALOGS, schemata, tables, and columns, and a match used to "
        "repeat that walk once per criterion.",
    )

    heading(doc, "3.2 Cache metadata for a short time", 2)
    para(
        doc,
        "Caffeine is already a backend dependency. Cache catalog names, schema names, table "
        "names, and column name plus type. Key the column cache by catalog, schema, and table. "
        "Use a short time-to-live, on the order of one to five minutes, so a newly registered "
        "table appears without a restart. Drop the cache entry when a registration or column "
        "override for that table is saved or deleted. The Query Builder cascade and SQL preview "
        "then share one lookup instead of each issuing information_schema queries.",
    )

    heading(doc, "3.3 Stop the cascade from refetching the same lists", 2)
    para(
        doc,
        "In LakehouseCascade, run the catalog effect from the browse-filter key, the schema "
        "effect from catalog plus filter, and the table effect from catalog, schema, and filter. "
        "Do not list value.catalog, value.schema, or onChange as reasons to start a new request "
        "when those values did not change the query. After the list returns, call onChange only "
        "when the current selection is absent from the list. If the selection is still present, "
        "leave the parent state alone.",
    )

    heading(doc, "3.4 Preview SQL only when the question has settled", 2)
    para(
        doc,
        "ExtractRecordsTab should not post /api/analysis/match.sql on every keystroke. Wait "
        "on the order of one second after the last change, and only when the source table and "
        "the display columns are complete. Pass an AbortSignal into the fetch and abort the "
        "previous preview when the extract changes or the component unmounts. The backend "
        "should treat a closed client as a reason to stop the metadata query rather than hold "
        "the Presto connection until the driver timeout.",
    )

    heading(doc, "3.5 Separate metadata checks from the match engine", 2)
    para(
        doc,
        "A registered table already has a catalog, schema, table, and column list the product "
        "trusts for Query Builder. Preview and match validation should use that registration "
        "and the cache. A live information_schema walk belongs on an explicit refresh, not on "
        "every preview. The match itself stays on the existing AnalysisExecutionRouter: one "
        "business database and plain SQL go direct JDBC; a lakehouse table, a fuzzy or "
        "folded-group question, or more than one source goes to Presto, and only when each "
        "business source has a catalog alias.",
    )

    heading(doc, "3.6 Timeouts and the connection pool", 2)
    para(
        doc,
        "Do not enlarge the Presto pool until connections are held for less time. The log "
        "showed slow checkouts that all returned, not abandoned connections. Full match "
        "statements keep SRSE_QUERY_TIMEOUT_SECONDS, which defaults to 180. Metadata reads "
        "need a much shorter timeout, on the order of 10 to 15 seconds, so a stuck DB2 "
        "information_schema call cannot occupy a pool slot for the whole match budget.",
    )
    table(
        doc,
        ["Change", "Where", "Expected effect"],
        [
            ["One validateColumns call per table", "RecordMatchService", "One metadata walk per table instead of one per column"],
            ["Short-lived metadata cache", "Lakehouse browse / registry", "Cascade and preview stop repeating information_schema"],
            ["Stable cascade dependencies", "LakehouseCascade.tsx", "The 42 repeated catalog calls stop"],
            ["Longer debounce and abort", "ExtractRecordsTab.tsx", "Stale previews do not keep Presto connections"],
            ["Cached columns for validation", "Registry describe path", "Preview returns SQL without a live catalog walk"],
            ["Shorter metadata timeout", "Browse JDBC calls", "A stuck metadata query releases the pool slot"],
        ],
    )

    heading(doc, "4. Exception handling plan")
    para(
        doc,
        "Two controller-advice classes already map failures. DataAccessExceptionHandler turns "
        "a lost Presto or database connection into HTTP 503 and other SQL failures into HTTP 500, "
        "including the driver message and, for UncategorizedSQLException and BadSqlGrammarException, "
        "the failing SQL. ApiExceptionHandler maps IllegalArgumentException and a bad saved-query "
        "payload to 400, AdminAccessDeniedException to 403, IllegalStateException to 409, and a "
        "failed audit write to 503. Streaming match errors never reach these classes. The status "
        "line is already sent, and RecordMatchService writes an error line into the NDJSON body.",
    )

    heading(doc, "4.1 One error shape for synchronous APIs", 2)
    para(
        doc,
        "Return a small JSON object with code, message, and requestId. Keep the message in "
        "the language an officer can act on. Do not return a stack trace or a raw exception "
        "class name. The existing string bodies can stay readable inside the message field so "
        "current screens that show the response text still have a sentence to display. New "
        "clients should read the code.",
    )
    table(
        doc,
        ["Situation", "HTTP status", "code"],
        [
            ["Validation, unknown table or column, bad predicate", "400", "invalid_request"],
            ["Authenticated caller lacks the role", "403", "forbidden"],
            ["Conflicting state, such as a missing required setup", "409", "conflict"],
            ["Presto or the business database cannot be reached", "503", "connection_unavailable"],
            ["The statement ran and the engine rejected it", "500", "query_failed"],
            ["Audit row could not be written before an export", "503", "audit_unavailable"],
        ],
    )

    heading(doc, "4.2 Classify connectivity narrowly", 2)
    para(
        doc,
        "DataAccessExceptionHandler treats any IOException in the cause chain as a lakehouse "
        "outage. That was added because the Presto driver wraps a dropped socket in a generic "
        "IOException. Limit the 503 path to socket, connect, timeout, and unknown-host failures, "
        "plus the Presto transport IOException only when the message or the SQL state shows a "
        "connection problem. A query that ran and failed stays 500 and still includes the "
        "failing SQL, which is parameterised and does not contain officer-typed values.",
    )

    heading(doc, "4.3 Streaming match uses the same messages", 2)
    para(
        doc,
        "POST /api/analysis/match keeps its NDJSON error line, because the HTTP status can no "
        "longer change. Build that line with the same distinction: connection unavailable, "
        "validation, or query failed. QueryExecutionMonitor already records FAILURE versus "
        "CANCELLED when the client closes the stream. A client disconnect should cancel the "
        "Presto statement instead of leaving it in the pool.",
    )

    heading(doc, "4.4 Authentication failures and health checks", 2)
    para(
        doc,
        "SessionBearerAuthenticationFilter catches every RuntimeException and leaves the "
        "request anonymous. An invalid token should stay anonymous. Any other failure should "
        "be logged with the request path before the request continues, so a parser bug is not "
        "silent. HealthController and ConnectionInfoController catch Exception and report the "
        "plane as down. Catch the data-access and connection exceptions there. Let a programming "
        "error surface as a server error instead of a false 'plane down'.",
    )

    heading(doc, "4.5 What the browser shows", 2)
    bullet(doc, "HTTP 503: tell the officer the database connection is unavailable and that they can retry. Point them at the admin connection screen only when the caller is an administrator.")
    bullet(doc, "HTTP 400: show the validation message from the API. Do not replace it with a generic failure.")
    bullet(doc, "HTTP 500 on a query: show the driver message the API already returns, including which statement failed when the body includes it.")
    bullet(doc, "Abort a superseded match.sql call so an older 400 or 503 cannot replace the preview for the question now on screen.")

    heading(doc, "5. Code standardization")
    para(
        doc,
        "These rules follow structures already in the repository. They are not a new framework.",
    )
    bullet(doc, "Column checks for several columns of one table go through validateColumns or describeColumns. validateColumn stays for a single column. New match-planning code does not loop validateColumn.")
    bullet(doc, "Synchronous API errors go through DataAccessExceptionHandler and ApiExceptionHandler. A controller does not invent its own status and string body for the same cases.")
    bullet(doc, "Frontend loads use a cancelled flag and AbortController. Effect dependencies are the values that change the request. An effect does not call setState in a way that immediately starts the same request again. ExtractRecordsTab and RuleColumnValuePicker already trip the set-state-in-effect lint rule; fix those two call sites as cleanup, not as a rewrite of the screens.")
    bullet(doc, "SQL sent to Presto or to a saved JDBC connection stays parameterised. Catalog, schema, table, and column identifiers stay on the LakehouseIdentifiers allow-list before they are interpolated.")
    bullet(doc, "Timeouts, cohort cap, and estimated-row cap stay in application.yml (SRSE_QUERY_TIMEOUT_SECONDS and the existing SRSE_ANALYSIS_* and SRSE_COHORT_CAP settings). Services do not add a second hardcoded timeout.")
    bullet(doc, "Java stays on 17 with the current Maven layout. The frontend stays on the current ESLint setup. Do not introduce a second formatter or a second error-response library as part of this work.")

    heading(doc, "6. Implementation order")
    table(
        doc,
        ["Step", "Work", "Done when"],
        [
            ["1", "Batch column validation in RecordMatchService", "One SQL preview issues one metadata walk per table, visible in a single-column versus multi-column timing check"],
            ["2", "Cache browse results and invalidate on registration changes", "Repeating the catalog cascade does not open a new information_schema query inside the TTL"],
            ["3", "Stabilize LakehouseCascade and abort stale match.sql calls", "A quiet Query Builder stops issuing catalog GETs, and changing the question cancels the previous preview"],
            ["4", "Shorter timeout for metadata only", "A hung information_schema call returns inside the metadata budget and releases the Presto connection"],
            ["5", "Shared JSON error body and narrower 503 classification", "Synchronous failures return code, message, and requestId; a non-connection IOException is not labeled lakehouse unavailable"],
            ["6", "Same messages on the match stream, and logging for unexpected token-parse failures", "An NDJSON error line uses the same wording as the synchronous handler"],
        ],
    )
    para(
        doc,
        "Steps 1 to 3 are the ones that address the 14 October log. Steps 4 to 6 keep the "
        "same failure from filling the pool again and make the message the officer sees match "
        "the failure. Pool size stays unchanged until a repeat of that log shows connections "
        "returning in seconds rather than minutes.",
    )

    heading(doc, "7. Out of scope")
    bullet(doc, "No change to role rules, scope deny-by-default, or the Presto-versus-direct-JDBC choice for a federated question.")
    bullet(doc, "No change to the operational PostgreSQL schema.")
    bullet(doc, "This document does not itself modify application code.")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
