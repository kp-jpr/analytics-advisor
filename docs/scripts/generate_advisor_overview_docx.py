"""Build docs/Analytics_Advisor_Overview_Architecture_and_Roadmap.docx."""
from __future__ import annotations

from datetime import date
from pathlib import Path

import matplotlib.pyplot as plt
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch

ROOT = Path(__file__).resolve().parent.parent
ASSETS = ROOT / "tech_doc_assets"
OUT = ROOT / "Analytics_Advisor_Overview_Architecture_and_Roadmap.docx"


def _box(ax, x, y, w, h, text, fc="#E8EEF7", ec="#1E2761", fs=8.5):
    ax.add_patch(
        FancyBboxPatch(
            (x, y),
            w,
            h,
            boxstyle="round,pad=0.02,rounding_size=0.08",
            linewidth=1.2,
            edgecolor=ec,
            facecolor=fc,
        )
    )
    ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=fs)


def _arrow(ax, x1, y1, x2, y2):
    ax.add_patch(
        FancyArrowPatch(
            (x1, y1),
            (x2, y2),
            arrowstyle="-|>",
            mutation_scale=11,
            linewidth=1.1,
            color="#444444",
        )
    )


def _save(fig, name: str) -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    fig.savefig(ASSETS / name, dpi=160, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def diagram_role_menus() -> None:
    fig, ax = plt.subplots(figsize=(11.2, 6.4))
    ax.set_xlim(0, 11.2)
    ax.set_ylim(0, 6.4)
    ax.axis("off")
    ax.set_title("Menus visible to each role", fontsize=14, fontweight="bold", color="#1E2761")
    _box(ax, 0.3, 3.3, 2.5, 2.5, "Officer\n\nOverview\nQuery Builder\nDashboard", fc="#E8EEF7")
    _box(
        ax,
        3.1,
        2.2,
        2.6,
        3.6,
        "Admin\n\nOfficer screens\nUsers and Roles\nScope Hierarchy\nConnections\nTable Registry\nColumn Settings\nBackup\nGuardrails",
        fc="#FFFFFF",
        fs=8,
    )
    _box(
        ax,
        6.0,
        1.3,
        2.5,
        4.5,
        "Super Admin\n\nEverything Admin sees\nData Sources\nOTP Gateways\nAudit Log",
        fc="#CADCFC",
        fs=8,
    )
    _box(
        ax,
        8.8,
        3.6,
        2.1,
        2.2,
        "Audit Reader\nalone\n\nNo Admin menu\nAudit API only",
        fc="#FFF4E5",
        fs=8,
    )
    _save(fig, "overview_role_menus.png")


def diagram_two_planes() -> None:
    fig, ax = plt.subplots(figsize=(11.2, 6.6))
    ax.set_xlim(0, 11.2)
    ax.set_ylim(0, 6.6)
    ax.axis("off")
    ax.set_title("Two data planes plus department databases", fontsize=14, fontweight="bold", color="#1E2761")
    _box(ax, 4.2, 5.4, 2.8, 0.8, "Browser\nOfficer or Super Admin", fc="#CADCFC")
    _box(ax, 0.4, 3.6, 3.2, 1.2, "Frontend\nNext.js 16 / React 19\nlocalhost:3000", fc="#FFFFFF")
    _box(ax, 4.0, 3.6, 3.2, 1.2, "Backend\nJava 17 / Spring Boot 3.3\nlocalhost:8080", fc="#FFFFFF")
    _box(ax, 7.8, 3.6, 3.0, 1.2, "Data Sources\nExtra PostgreSQL or DB2\nregistered tables", fc="#FFFFFF")
    _box(ax, 0.4, 1.2, 3.2, 1.5, "Operational plane\nPostgreSQL :5433\nUsers, roles, scopes, audit", fc="#FFF4E5")
    _box(ax, 4.0, 1.2, 3.2, 1.5, "Analytical plane\nPrestoDB 0.297\nIceberg lakehouse SQL", fc="#FFF4E5")
    _arrow(ax, 5.6, 5.4, 2.0, 4.8)
    _arrow(ax, 5.6, 5.4, 5.6, 4.8)
    _arrow(ax, 3.6, 4.2, 4.0, 4.2)
    _arrow(ax, 5.2, 3.6, 2.0, 2.7)
    _arrow(ax, 5.6, 3.6, 5.6, 2.7)
    _arrow(ax, 7.2, 4.2, 7.8, 4.2)
    _save(fig, "overview_two_planes.png")


def diagram_query_route() -> None:
    fig, ax = plt.subplots(figsize=(11.2, 6.2))
    ax.set_xlim(0, 11.2)
    ax.set_ylim(0, 6.2)
    ax.axis("off")
    ax.set_title("How a query chooses its database", fontsize=14, fontweight="bold", color="#1E2761")
    _box(ax, 3.8, 4.8, 3.6, 0.9, "Officer runs a query\nin Query Builder", fc="#CADCFC")
    _box(ax, 3.6, 3.2, 4.0, 1.0, "AnalysisExecutionRouter\nlooks at tables and SQL", fc="#FFFFFF")
    _box(ax, 0.4, 1.0, 4.6, 1.4, "Direct JDBC\nOne external table\nPlain SQL, no fuzzy match\nGoes to that database", fc="#E8F5E9")
    _box(ax, 6.2, 1.0, 4.6, 1.4, "Presto\nLakehouse, fuzzy, or\nmore than one source\nNeeds a federation alias", fc="#FFF4E5")
    _arrow(ax, 5.6, 4.8, 5.6, 4.2)
    _arrow(ax, 4.2, 3.2, 2.6, 2.4)
    _arrow(ax, 7.0, 3.2, 8.4, 2.4)
    _save(fig, "overview_query_route.png")


def heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def para(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(11)


def bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text, style="List Bullet")
    for r in p.runs:
        r.font.name = "Arial"
        r.font.size = Pt(11)


def table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = h
        for p in t.rows[0].cells[i].paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.name = "Arial"
                r.font.size = Pt(9)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            t.rows[ri + 1].cells[ci].text = val
            for p in t.rows[ri + 1].cells[ci].paragraphs:
                for r in p.runs:
                    r.font.name = "Arial"
                    r.font.size = Pt(9)
    doc.add_paragraph()


def figure(doc: Document, png: str, caption: str) -> None:
    path = ASSETS / png
    doc.add_picture(str(path), width=Inches(6.4))
    c = doc.add_paragraph(caption)
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in c.runs:
        r.italic = True
        r.font.size = Pt(10)
        r.font.name = "Arial"
    doc.add_paragraph()


def build() -> None:
    diagram_role_menus()
    diagram_two_planes()
    diagram_query_route()

    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Analytics Advisor")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(0x1E, 0x27, 0x61)
    run.font.name = "Arial"

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run(
        "Purpose, features, menus by role,\ntechnical architecture, limitations, and improvement plan"
    )
    sr.font.size = Pt(14)
    sr.font.name = "Arial"

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mr = meta.add_run(
        f"Version 1.0  |  {date.today():%d %B %Y}\n"
        "Product name: Analytics Advisor (SRSE — Scheme Rule Simulation Engine)\n"
        "This document describes the running product. It does not change application code."
    )
    mr.font.size = Pt(10)
    mr.font.name = "Arial"
    doc.add_page_break()

    heading(doc, "1. Purpose", 1)
    para(
        doc,
        "Analytics Advisor lets a departmental officer ask who qualifies for a welfare scheme "
        "without writing SQL. The officer picks tables, join keys, rules, and thresholds. "
        "The product compiles that choice into a parameterised query, runs it, and shows the "
        "matching people. It also records who ran the query and hides rows that fall outside "
        "the officer’s geography.",
    )
    para(
        doc,
        "The same product is used by administrators to create accounts, attach department "
        "databases, and register the tables officers are allowed to see. Officers never edit "
        "those connections. Super Admins do.",
    )
    para(
        doc,
        "Example. An officer in Jaipur wants people in the pension table whose age is at least "
        "18 and whose name is similar to a name in the lakehouse beneficiary table. They choose "
        "both tables, set a fuzzy match on the name, and run the query. The product returns "
        "only Jaipur rows when that officer’s scope is Jaipur, and it writes an audit event "
        "for the run.",
    )

    heading(doc, "2. Features", 1)
    para(doc, "Each feature below is present in the current application.")

    heading(doc, "2.1 Sign-in and language", 2)
    bullet(doc, "Local sign-in with a username and a BCrypt password. A new account can be required to change the password before the rest of the application opens.")
    bullet(doc, "Five failed attempts lock the account for 15 minutes. During the lock, the screen shows the same sentence as a wrong password.")
    bullet(doc, "A session token lasts about eight hours. English and Hindi are available from the sidebar.")
    bullet(doc, "Example. superadmin signs in, is sent to change the password on first use, then opens Users & Roles.")

    heading(doc, "2.2 Officer workspace", 2)
    bullet(doc, "Overview shows the registered data the officer is allowed to browse.")
    bullet(doc, "Query Builder is where a query is composed: tables, join keys, rules, grouping, aggregates, deduplication, and mismatch-only results.")
    bullet(doc, "Dashboard shows the result of a completed run and can export it.")
    bullet(doc, "Exact joins and fuzzy name joins are both supported. A fuzzy join compares similar text, such as two spellings of the same name.")
    bullet(doc, "Count-only runs return a row count without listing every person. Multi-target match compares one source table with several target tables in one request.")
    bullet(doc, "Results can be exported as CSV, JSON, XML, or Excel. A query definition can be saved and opened again.")

    heading(doc, "2.3 Administration", 2)
    bullet(doc, "Users and roles, scope hierarchy, the two application connections, external data sources, table registry, and column settings.")
    bullet(doc, "Scope is deny-by-default. An officer with no geography binding sees no beneficiary rows.")
    bullet(doc, "Audit records sign-in, query, and administration events. OTP gateways send one-time passwords when multi-factor sign-in is turned on for an account.")
    bullet(doc, "Configuration backup downloads connections (without passwords), registrations, and column settings. Analysis guardrails show the deployment limits and cannot be edited on screen.")
    bullet(doc, "Column metadata is snapshotted for the pickers. Live column types are still read when SQL is compiled.")

    heading(doc, "2.4 Guardrails an officer will meet", 2)
    table(
        doc,
        ["Limit", "What it does", "Example"],
        [
            ["Query timeout", "Stops a match that runs too long. Default 180 seconds.", "A fuzzy match on a very large table is cancelled and the officer sees a failure message."],
            ["Result cap", "Limits how many rows a run may return.", "The on-screen preview stops at the configured cap. A full export is a separate download."],
            ["Fan-out ceiling", "Refuses a join that would multiply into too many rows.", "Matching every person to every village without a selective key is refused before the full join runs."],
            ["Scope", "Adds the officer’s geography to the query.", "A Jaipur officer does not receive rows for another district."],
        ],
    )

    heading(doc, "3. Menus by role", 1)
    para(
        doc,
        "The product has four roles: Super Admin, Admin, Officer, and Audit Reader. "
        "The sidebar the officer sees is Overview, Query Builder, and Dashboard. "
        "The Admin link appears only when the signed-in account has the administrator authority. "
        "Inside Admin, three groups appear: Access Management, Data Configuration, and Security & Operations.",
    )
    figure(
        doc,
        "overview_role_menus.png",
        "Figure 1. Officer, Admin, Super Admin, and Audit Reader do not see the same menus.",
    )

    heading(doc, "3.1 Officer", 2)
    bullet(doc, "Overview, Query Builder, and Dashboard.")
    bullet(doc, "Theme and language controls sit at the bottom of the sidebar.")
    bullet(doc, "Opening the admin address sends the officer back to Overview.")
    para(
        doc,
        "Example. An officer opens Query Builder, selects the registered beneficiary table, "
        "filters district to the values their scope allows, and runs the query. They cannot "
        "open Connections or Data Sources.",
    )

    heading(doc, "3.2 Admin", 2)
    bullet(doc, "The three officer screens, plus the Admin link.")
    bullet(doc, "Access Management: Users & Roles, Scope Hierarchy.")
    bullet(doc, "Data Configuration: Connections, Table Registry, Column Settings.")
    bullet(doc, "Security & Operations: Configuration Backup and Analysis Guardrails.")
    bullet(doc, "Data Sources and OTP Gateways stay hidden. Audit Log stays hidden unless this account also has the Audit Reader role.")
    para(
        doc,
        "Example. An admin creates jaipur.officer, assigns the Officer role, and binds the "
        "Jaipur scope node. That officer then sees only Jaipur rows. The admin cannot add a "
        "new department database; that action is reserved for Super Admin.",
    )

    heading(doc, "3.3 Super Admin", 2)
    bullet(doc, "Every Admin item, plus Data Sources and OTP Gateways.")
    bullet(doc, "Audit Log is included because a Super Admin is treated as an audit reader.")
    bullet(doc, "Example. Super Admin adds Pension DB on Data Sources, tests it, registers public.beneficiary, then sets business names for its columns.")

    heading(doc, "3.4 Audit Reader", 2)
    para(
        doc,
        "The Audit Reader role grants the audit-read authority and nothing else. "
        "The admin page itself requires the administrator flag. An account that is only "
        "an Audit Reader therefore does not see the Admin menu and is sent to Overview "
        "if they open the admin address. The audit API still expects that authority. "
        "This gap is listed again under limitations.",
    )

    heading(doc, "3.5 What each admin menu item is for", 2)
    table(
        doc,
        ["Menu", "Who sees it", "What it is for", "Example"],
        [
            ["Users & Roles", "Admin and Super Admin", "Create accounts, assign roles, activate or deactivate users, reset passwords.", "Create jaipur.officer with role Officer."],
            ["Scope Hierarchy", "Admin and Super Admin", "Maintain the geography tree that limits which rows an officer sees.", "Add District Jaipur and assign it to that officer."],
            ["Connections", "Admin and Super Admin", "The two databases the application itself uses: operational store and Presto.", "Point the operational card at PostgreSQL on port 5433. The analytical card shows whether Presto answers SELECT 1."],
            ["Data Sources", "Super Admin", "Add a department PostgreSQL or DB2 database, browse its tables, and register the ones officers may query.", "Add Pension DB and register public.beneficiary."],
            ["Table Registry", "Admin and Super Admin", "Register lakehouse tables and the display tags officers see in analysis.", "Register iceberg.srse.beneficiary with the tag Beneficiary."],
            ["Column Settings", "Admin and Super Admin", "Business names, fuzzy matching, visibility, and how two columns are compared.", "Mark name as a fuzzy column and hide an internal key."],
            ["Audit Log", "Super Admin, or an admin who also has Audit Reader", "Review sign-in, query, and administration events.", "See that jaipur.officer ran a match at 11:00."],
            ["OTP Gateways", "Super Admin", "Email and SMS settings used when an account requires a one-time password.", "Point email at the department SMTP server."],
            ["Configuration Backup", "Admin and Super Admin", "Download or restore connections, registrations, and column settings. Passwords are not in the file.", "Export before changing the Presto URL, then re-enter the password on Connections."],
            ["Analysis Guardrails", "Admin and Super Admin", "Read-only view of match limits set by deployment configuration.", "See the maximum estimated fan-out and the multi-match time budget."],
        ],
    )

    heading(doc, "3.6 Connections and Data Sources", 2)
    para(
        doc,
        "Connections always has two cards. Operational holds users, roles, scopes, audit, "
        "and saved configuration. On a local machine that is PostgreSQL on port 5433, database srse. "
        "A saved operational change is tested and then takes effect after the backend restarts. "
        "Analytical is the Presto query engine. A saved analytical change is tested and applied immediately. "
        "The address jdbc:presto://host:8080 is enough. A trailing catalog and schema is only a default, "
        "because every table is named catalog.schema.table.",
    )
    para(
        doc,
        "Data Sources is a list a Super Admin grows. Each entry is a department database with its own "
        "host, port, username, and password. After a successful test, the explorer lists schemas, tables, "
        "views, and columns. Registering a table makes it available in analysis. A PostgreSQL source can "
        "also receive a Presto federation alias so it can be joined with the lakehouse. A DB2 source can "
        "answer a query that stays on that one database. This Presto build has no DB2 connector, so a DB2 "
        "source cannot be joined to another database.",
    )
    para(
        doc,
        "Example of using both. Connections keeps superadmin on the operational database and keeps Presto "
        "as the analytical engine. Data Sources adds pension-db on port 5432. An officer filters that one "
        "table with an exact district match: the SQL goes straight to the pension database. The same officer "
        "fuzzy-matches names against the lakehouse table: that query goes through Presto, using the federation alias.",
    )

    heading(doc, "4. Technical architecture", 1)
    para(
        doc,
        "The browser talks to a Next.js 16 and React 19 frontend. The frontend talks to a Java 17 "
        "Spring Boot 3.3 backend. The backend has two permanent data planes and an optional list of "
        "external JDBC sources.",
    )
    figure(
        doc,
        "overview_two_planes.png",
        "Figure 2. Frontend, backend, operational store, analytical engine, and department data sources.",
    )
    table(
        doc,
        ["Part", "Technology", "What it holds"],
        [
            ["Frontend", "Next.js 16.1, React 19, TypeScript, HeroUI, Tailwind", "Screens, session token in the browser, calls to the backend."],
            ["Backend", "Java 17, Spring Boot 3.3, Spring Security", "Query compilation, execution, users, audit, and admin APIs."],
            ["Operational plane", "PostgreSQL 16 locally (DB2 is the other supported operational engine), JPA, Liquibase", "Users, roles, scopes, audit events, saved queries, registrations, column metadata."],
            ["Analytical plane", "PrestoDB 0.297 JDBC, Iceberg tables", "Lakehouse SQL. Local stack also uses Hive metastore and MinIO."],
            ["Data sources", "PostgreSQL or IBM DB2 JDBC, password encrypted at rest", "Department databases whose tables have been registered."],
            ["Sign-in", "BCrypt passwords, HS256 JWT, about eight hours", "Local mode. Mock mode and a RajSewadwar mode also exist in configuration."],
        ],
    )

    heading(doc, "4.1 How a query is executed", 2)
    para(
        doc,
        "The officer’s selections are compiled into parameterised SQL. Values the officer typed "
        "are bound as parameters. The router then chooses where that SQL runs.",
    )
    figure(
        doc,
        "overview_query_route.png",
        "Figure 3. One plain query on one external table goes direct. Lakehouse, fuzzy, and multi-source queries go to Presto.",
    )
    bullet(doc, "No external table, or a lakehouse table: the SQL runs on Presto.")
    bullet(doc, "One registered external table and plain SQL, with no fuzzy match and no Presto-only construct: the SQL is rewritten for that database and sent over JDBC.")
    bullet(doc, "Fuzzy matching, a Presto-only construct, or more than one source: the SQL runs on Presto. Each external source in that query must have a federation catalog alias. DB2 cannot have one.")
    bullet(doc, "Before a large join, a fan-out check estimates the row count and refuses the query when the estimate is above the ceiling.")
    bullet(doc, "If the browser disconnects during a streamed result, the running statement is cancelled.")
    bullet(doc, "Failures that officers see are returned as a structured error with a code and a message, rather than a raw driver stack.")

    heading(doc, "4.2 Request path", 2)
    bullet(doc, "The browser stores the session token and sends it to the backend.")
    bullet(doc, "The backend checks the token, the password-change flag, and the session version. A password change invalidates older tokens.")
    bullet(doc, "Scope predicates are added so the officer’s geography is part of the SQL.")
    bullet(doc, "Results stream back for the on-screen table. A full export is a separate download and stops if the file write fails. It does not write an error into the middle of the file.")
    bullet(doc, "An audit event records the action, the outcome, and, for queries, the shape of the SQL and the duration.")

    heading(doc, "5. Limitations", 1)
    para(
        doc,
        "These are limits of the current release. They are not defects in the menus described above. "
        "Each one has a consequence an administrator or an officer can actually meet.",
    )
    table(
        doc,
        ["Limitation", "What the user sees", "Why"],
        [
            ["Presto must be running", "The Analytical card reports a SELECT 1 failure. Lakehouse, fuzzy, and cross-database queries fail.", "The analytical plane is a separate process. A laptop without Docker does not start Presto, MinIO, and the Hive metastore."],
            ["DB2 cannot join other databases", "A DB2 data source is marked single-source only. Fuzzy match and lakehouse joins are refused.", "PrestoDB 0.297 ships no DB2 connector, so no federation alias can exist."],
            ["Oracle, MySQL, MariaDB, and SQL Server cannot be added", "The Add connection list offers PostgreSQL and IBM DB2. Creating another type is refused.", "Those types are defined, including an Oracle thin URL and the Oracle driver, and they are turned off for this release."],
            ["Operational edits wait for a restart", "Connections says the operational save applies after a backend restart.", "The operational pool is the application store. The analytical pool can be swapped while the process is running."],
            ["External JDBC is not pooled", "Repeated queries to the same department database open connections as needed.", "Each data source uses a direct connection rather than a standing pool."],
            ["Audit Reader cannot open Admin", "An account with only the Audit Reader role has no Admin link and is redirected away from the admin page.", "The admin page requires the administrator authority. Audit read alone does not grant it."],
            ["Lockout looks like a bad password", "After five failures the message is still “Invalid username or password” for 15 minutes.", "The sign-in response is the same for an unknown user, a bad password, and a lock, so the screen does not reveal which account exists."],
            ["Fan-out and comparison repeat work", "A very wide join is estimated with a live count. A comparison summary runs the query again.", "The estimate and the summary are separate executions. They are not served from a stored result."],
            ["RajSewadwar sign-in is not active", "Department single sign-on is a configured mode and does not authenticate a user yet.", "The integration point exists. The live identity payload is not accepted."],
        ],
    )

    heading(doc, "6. Plan to overcome the limitations", 1)
    para(
        doc,
        "The phases below are a recommended order of work. This document does not implement them. "
        "Phase 1 makes the product usable for lakehouse queries. Phase 2 widens which databases "
        "and which roles can do their job. Phase 3 improves sign-in and the cost of large queries.",
    )

    heading(doc, "6.1 Phase 1 — Make the analytical engine available", 2)
    bullet(doc, "Run PrestoDB 0.297 with the Hive metastore and MinIO, using the project’s Docker compose services, on a machine with enough free disk.")
    bullet(doc, "Confirm the Analytical card returns up, which is SELECT 1 against jdbc:presto://localhost:8081.")
    bullet(doc, "Replace the raw driver sentence on that card with a short statement that the analytical engine is not running, and keep the technical detail for the administrator.")
    bullet(doc, "Outcome. Lakehouse matches, fuzzy matches, and federated joins can run. A single-table plain query on a data source already can run without Presto; this phase does not change that path.")

    heading(doc, "6.2 Phase 2 — Widen data sources and the audit role", 2)
    bullet(doc, "Enable Oracle on Data Sources. The type, port 1521, thin URL, and JDBC driver are already in the backend. The work is to turn the type on, verify browse, register, direct SQL, and a Presto Oracle catalog alias.")
    bullet(doc, "Give each saved external database a small connection pool so repeated officer queries reuse connections.")
    bullet(doc, "Let an Audit Reader open the Audit Log without granting Users & Roles, Connections, or Data Sources. The audit API already checks the audit-read authority.")
    bullet(doc, "Outcome. A department Oracle database can be registered the same way Pension DB is registered today. An auditor can read the log.")

    heading(doc, "6.3 Phase 3 — Clearer sign-in and less repeated work", 2)
    bullet(doc, "Keep the public sign-in message the same for unknown users, and show a distinct lock message only after the username is known to be real, so officers are not locked out without an explanation.")
    bullet(doc, "Finish RajSewadwar so a departmental identity can open a session with the same roles and scopes as a local account.")
    bullet(doc, "Reuse the fan-out estimate and avoid running a full comparison query a second time when the match result is already in hand.")
    bullet(doc, "Outcome. Sign-in failures are understandable, departmental single sign-on works, and large comparisons do less duplicate work.")

    heading(doc, "7. What this document does not change", 1)
    para(
        doc,
        "No screen, API, or database setting is modified by this document. Connections remains "
        "the two application planes. Data Sources remains the list of department databases. "
        "Oracle remains unavailable until Phase 2 is carried out in the product.",
    )

    doc.save(OUT)
    print(f"wrote {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    build()
